"""Build the HW1 write-up (Word) from the verified results (hw1/results_log.md) and figures.

All numbers come from hw1/results_log.md (outputs verified on the Yale cluster). The one chart
drawn here (combined-diabetes chronic-condition chart) re-plots logged values; no data are re-run.
Run with:  uv run --with python-docx --with matplotlib python build_writeup.py
"""
import re
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HW1 = r"C:\Users\mnich\OneDrive\Documents\Second Year\EmpiricalStrategy\hw1"
FIG = HW1 + r"\figures\\"
REPORT = HW1 + r"\report\\"
OUT = sys.argv[1] if len(sys.argv) > 1 else REPORT + "MGT634_HW1_NHAMCS_writeup.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x52, 0x51, 0x4E)

# ---------------------------------------------------------------- report-only chart
# Chronic-condition prevalence with the three diabetes checkboxes combined, so the chart
# matches the text. Values: results_log.md, Q6 (weighted % of visits with the section completed).
chronic = [("Hypertension", 23.97), ("Any diabetes (type 1, 2 or unspecified)", 11.06),
           ("Asthma", 9.98), ("Depression", 9.45), ("Hyperlipidemia", 8.17),
           ("Substance abuse", 6.65), ("Coronary artery disease", 6.08), ("COPD", 5.37),
           ("Obesity", 3.64), ("Cancer", 3.40)]
fig, ax = plt.subplots(figsize=(8, 3.0))
names, vals = [c for c, _ in chronic][::-1], [v for _, v in chronic][::-1]
bars = ax.barh(names, vals, color="#2a78d6", height=0.65)
ax.bar_label(bars, fmt="%.1f%%", padding=3, fontsize=9, color="#52514e")
ax.set_xlabel("% of visits with the chronic-condition section completed (weighted)", color="#52514e")
for side in ["top", "right", "left"]:
    ax.spines[side].set_visible(False)
ax.tick_params(colors="#52514e")
ax.grid(axis="x", color="#e6e5e1", linewidth=0.8)
ax.set_axisbelow(True)
ax.set_xlim(0, 27)
fig.tight_layout()
CHRONIC_PNG = REPORT + "fig_chronic_combined_diabetes.png"
fig.savefig(CHRONIC_PNG, dpi=200)
plt.close(fig)

# ---------------------------------------------------------------- document setup
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.75)
sec.top_margin = sec.bottom_margin = Inches(0.65)

base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(9.5)
base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
base.paragraph_format.space_after = Pt(4)
base.paragraph_format.line_spacing = 1.05
h = doc.styles["Heading 1"]
h.font.name, h.font.size, h.font.bold = "Calibri", Pt(12.5), True
h.font.color.rgb = NAVY
h.paragraph_format.space_before = Pt(11)
h.paragraph_format.space_after = Pt(3)
h.paragraph_format.keep_with_next = True


def runs(p, text, size=None, color=None):
    """Add text to paragraph p; **bold** segments are rendered bold."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            r = p.add_run(part)
            r.bold = i % 2 == 1
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
    return p


def para(text, size=None, color=None, align=None, after=None):
    p = runs(doc.add_paragraph(), text, size, color)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text):
    p = runs(doc.add_paragraph(style="List Bullet"), text)
    p.paragraph_format.space_after = Pt(2)
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(rows, widths, caption, left_cols=(0,), size=8.5):
    """rows[0] = header. Light grid, shaded header, compact font."""
    cap = para(caption, size=9, color=NAVY, after=2)
    cap.runs[0].bold = True
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Inches(widths[j])
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            runs(c.paragraphs[0], str(val), size=size)
            if j not in left_cols:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i < len(rows) - 1:
                c.paragraphs[0].paragraph_format.keep_with_next = True
            if i == 0:
                shade(c, "E8EEF5")
                for r in c.paragraphs[0].runs:
                    r.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return t


def figure(paths, width, caption):
    """One image, or several side by side in a borderless table, plus caption."""
    if len(paths) == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(paths[0], width=Inches(width))
    else:
        t = doc.add_table(rows=1, cols=len(paths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, f in enumerate(paths):
            c = t.cell(0, j)
            c.width = Inches(width + 0.05)
            cp = c.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(0)
            cp.add_run().add_picture(f, width=Inches(width))
    cap = para(caption, size=8.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=9)
    cap.runs[0].italic = True


# ---------------------------------------------------------------- title block
t = para("Exploring U.S. Emergency Department Visits: NHAMCS-ED 2015", size=16, color=NAVY, after=0)
t.runs[0].bold = True
para("MGT 634 Empirical Strategy with AI · Homework 1 · Michal Kozuchowski, Laila Lapins, "
     "Nick Giamalis, Raymond Chang, Sean Weller · September 2026", size=9, color=GREY, after=6)

para("**Summary.** Weighted to the national level, the 2015 NHAMCS emergency department sample "
     "represents about 136.9 million U.S. ED visits. Arrivals follow a stable daily pattern: they rise "
     "from 7am, stay near their peak from 10am to 8pm, and are heaviest on Mondays. The median visit "
     "lasts about two and a half hours, so time in the department matters for crowding as much as "
     "arrival volume. Medicaid is the most common expected payer, and nearly half of visits with a "
     "completed chronic-condition section record at least one chronic condition. Triage level is "
     "strongly associated with admission, but recorded median waits differ little across levels 2-5.")

# ---------------------------------------------------------------- 1. data (Q1-Q2)
doc.add_heading("1. Data, sample design and cleaning (Q1-Q2)", level=1)
para("The NCHS public-use file has 21,061 rows and 1,031 columns. Each row is one sampled ED visit "
     "(one Patient Record Form), not a patient, so the same person can appear more than once. Hospital "
     "code plus patient code uniquely identifies every row, and the visits come from 248 EDs (a median "
     "of 94 sampled visits each). The unweighted count is therefore 21,061 visits. Summing the patient "
     "visit weight PATWT gives 136,943,181 estimated U.S. ED visits in 2015, the total NCHS reports "
     "(codebook p.28) and well above 100 million. The weights are uneven (117 to 42,002; the 10% of "
     "rows with the largest weights carry 28.7% of weighted visits), and the sample does not "
     "mirror the country: the Northeast is 20.4% of rows but 17.3% of weighted visits, the South 33.3% "
     "versus 37.8%. We therefore weight every total, share and rate. We report weighted point estimates "
     "only; we did not compute design-based standard errors (strata CSTRATM, clusters CPSUM), so small "
     "differences should not be over-read.")
para("Missing values are mostly not stored as blanks. A standard pandas NaN count flags 600 columns, "
     "but those are largely unused drug-category text fields. NCHS records missing answers as negative "
     "codes: -9 (blank), -8 (unknown) and -7 (not applicable), as defined item by item in the codebook "
     "(p.37 onward). Counting -9 and -8, 118 variables are more than 5% missing and 56 more than 20%. "
     "Among the variables we use, unimputed ethnicity is 24.2% blank (so we use the imputed ETHIM), "
     "pain score 29.5%, triage level 22.2%, wait time 15.2% (plus 3.4% not applicable because the "
     "patient was not seen by a provider), expected payer 8.7% and length of visit 7.0%. Intent "
     "(INTENT15) and hospital discharge status (HDSTAT) are asked only for injury visits and admitted "
     "patients, respectively, so their high missing shares are by design.")
para("Missingness also varies with observed characteristics. Wait time is missing for 7-12% of visits "
     "at triage levels 1-5 but for 38-45% of visits where triage was blank, not performed or not offered, "
     "and it ranges from 9.8% of Midwest visits to 19.7% of Northeast visits. We cannot tell whether the "
     "missing waits are longer or shorter than recorded ones, so our wait-time results describe visits "
     "with a recorded wait and may not generalize to the rest.")
para("Before computing any statistic we converted -9/-8/-7 to missing, recoded yes/no items to 0/1, and "
     "converted arrival time from military time to an arrival hour (1.3% blank). Length of visit reaches "
     "5,732 minutes and waits reach 1,305 minutes, so we emphasize medians and report means both raw and "
     "winsorized at the 99.5th percentile (98 and 86 visits capped) rather than deleting records.")
para("We used two kinds of checks. Externally, Table 1 compares our estimates with figures published in "
     "the NCHS documentation; all agree to rounding except wait-time missingness (15.2% vs 15.7%). "
     "Internally, the total-chronic-conditions count equals the sum of the 22 condition checkboxes on "
     "every answered visit, the medication count equals the number of filled medication slots on every "
     "visit, the weighted mean, median and correlation functions reproduce hand-computed results on test "
     "data, and a final end-to-end re-run of the notebook reproduces 20 key values recorded during the "
     "question-by-question review (a reproducibility check, not an external validation).")
table([["Estimate", "Our result", "NCHS published", "Agreement"],
       ["Records in file", "21,061", "21,061 (p.115)", "Exact"],
       ["Weighted U.S. ED visits", "136,943,181", "136,943,181 (p.28)", "Exact"],
       ["Female share of visits (weighted)", "55.44%", "55.436% (p.115)", "Rounding"],
       ["Hispanic share of visits (weighted)", "16.49%", "16.493% (p.115)", "Rounding"],
       ["Medicaid as primary expected payer, all visits", "31.15%", "31.152% (p.116)", "Rounding"],
       ["Visits with no medications (weighted)", "20.94%", "20.943% (p.117)", "Rounding"],
       ["Weighted drug mentions", "340,550,921", "340,550,921 (p.118)", "Exact"],
       ["BUN/creatinine test ordered, all visits", "6.67%", "6.7% (p.4)", "Rounding"],
       ["Missing: length of visit / expected payer", "7.0% / 8.7%", "7.0% / 8.7% (p.15)", "Exact"],
       ["Missing: wait time", "15.2%", "15.7% (p.15)", "0.5 points lower"]],
      [2.75, 1.15, 1.45, 1.35], "Table 1. Selected estimates compared with figures published by NCHS")

# ---------------------------------------------------------------- 2. Q3
doc.add_heading("2. Who comes to the ED, and how long they stay (Q3)", level=1)
table([["Statistic", "Unweighted", "Weighted"],
       ["Age, mean / median (years; 93 = top-coded \"93+\")", "37.6 / 34", "37.0 / 34"],
       ["Female / Hispanic (imputed ETHIM)", "55.1% / 15.9%", "55.4% / 16.5%"],
       ["Wait to see a provider, median (n = 17,153 with recorded wait)", "19 min", "18 min"],
       ["Wait, mean (winsorized at 99.5th percentile)", "39.5 min", "38.7 min"],
       ["Length of visit, mean raw / winsorized (n = 19,581)", "221.9 / 216.5 min", "213.7 / 209.8 min"],
       ["Length of visit, median", "154 min", "154 min"]],
      [3.8, 1.45, 1.45], "Table 2. Descriptive statistics, unweighted and weighted")
para("The median visit is by a patient aged 34; 55% of visits are by women and 16.5% by Hispanic "
     "patients (weighted). Half of patients with a recorded wait saw a provider within 18 minutes, but the mean "
     "wait is roughly twice the median, so a minority of long waits pulls the average up. The median "
     "visit lasts 154 minutes (2.6 hours) and the mean about 214 minutes. That long right tail matters "
     "for capacity, because long stays occupy beds and staff time. As a rough illustration, Little's law "
     "(average patients present = arrival rate x average time in the department) implies about 55,700 "
     "patients in U.S. EDs at a typical moment, or roughly 12 per ED across the 4,820 EDs NCHS estimates "
     "(codebook p.120), assuming steady arrivals.")
para("Weighting changes these statistics only slightly: mean age falls by 0.5 years, mean length of "
     "visit by about 8 minutes, and children's share of visits rises from 18.6% to 19.8%; the medians "
     "do not change. Weighting matters less for these headline figures than for volumes and subgroups: "
     "unweighted counts do not measure national volume (21,061 records versus 136.9 million visits), "
     "and the sample's regional mix differs from the country's.")

# ---------------------------------------------------------------- 3. Q4
doc.add_heading("3. When demand arrives (Q4)", level=1)
figure([FIG + "q4_arrivals_by_hour.png", FIG + "q4_visits_by_weekday.png"], 3.35,
       "Figure 1. Weighted arrivals by hour, as a share of visits with a recorded arrival time (98.7% of "
       "records), and average weighted visits per day by weekday")
para("Because 2015 had 53 Thursdays and 52 of every other weekday, we compare days as average visits "
     "per day. Monday is the busiest day (432,000 visits per day) and Sunday the quietest (346,000), a "
     "25% difference, with volume declining from Monday through Thursday. Over the year that is 22.45 "
     "million weighted Monday visits versus 18.00 million on Sundays. By hour, arrivals are lowest "
     "between 3am and 6am (about 1.3-1.4% of daily arrivals per hour), climb from 7am, and stay near "
     "their peak from 10am to 8pm; the busiest hour is 6pm (6.3%, 8.59 million arrivals over 2015). The peak hour receives 4.7 "
     "times as many arrivals as the quietest, and 69% of patients arrive between 10am and 10pm.")
para("For staffing, this points to planning around a long daytime and evening plateau rather than a "
     "single rush. Because the median patient stays about 2.6 hours, the number of patients in the "
     "department probably peaks after arrivals do, so coverage decisions should consider occupancy, "
     "which NHAMCS does not measure directly. The Monday excess is consistent with demand deferred over "
     "the weekend, but these data cannot establish the reason.")

# ---------------------------------------------------------------- 4. Q5
doc.add_heading("4. Payer mix (Q5)", level=1)
para("PAYTYPER is the primary expected source of payment, chosen by an NCHS hierarchy when a visit "
     "lists several; it is not a record of revenue actually collected. Across all visits (weighted), "
     "Medicaid/CHIP is the expected payer for 31.2%, private insurance 27.6%, Medicare 17.7% and "
     "self-pay 9.0%, other 2.2%, no charge 0.8% and workers' compensation 0.7%; for a further 10.8% "
     "the payer is blank or unknown. Medicaid and Medicare together account for "
     "48.9% of all visits and 54.8% of the 89.2% of visits with a known expected payer (Medicaid 34.9%, "
     "private 30.9%, Medicare 19.9%, self-pay 10.1%; self-pay and no charge together 11.0%). Age drives "
     "much of this mix (Figure 2): Medicaid is the expected payer for 65% of visits by children, Medicare "
     "for 78-87% of visits by patients 65 and older, and self-pay peaks at 17.5% among 25-44-year-olds.")
figure([FIG + "q5_payer_mix_by_age.png"], 3.7,
       "Figure 2. Primary expected payer by age group, weighted % of visits with a known expected payer")
para("If public programs pay less per visit than private insurers, as is commonly the case, an ED whose "
     "mix resembles this national profile would earn less per visit than one serving more privately "
     "insured patients, and self-pay and no-charge visits carry a higher risk of unpaid bills. Because "
     "EDs must screen and stabilize patients regardless of ability to pay (EMTALA), hospitals have little "
     "control over this mix at the door. These data cannot measure the financial effect directly: actual "
     "revenue depends on contracts and collections that NHAMCS does not record, and 10.8% of visits lack "
     "a recorded payer.")

# ---------------------------------------------------------------- 5. Q6-Q9
doc.add_heading("5. Clinical profile (Q6-Q9)", level=1)
figure([CHRONIC_PNG], 4.6,
       "Figure 3. Hypertension and diabetes are the most commonly recorded chronic conditions: weighted % of visits with the chronic-condition "
       "section completed (98.4% of records). \"Any diabetes\" combines three checkboxes; type 2 alone is "
       "4.7%.")
para("**Chronic conditions (Q6).** Among visits with the chronic-condition section completed (1.6% of "
     "records were blank), 47.6% record at least one chronic condition (weighted; 48.5% unweighted) and "
     "14.9% record three or more. Hypertension is the most common (24.0%). Diabetes is split across three "
     "checkboxes (type 1, type 2 and unspecified); counting any of them, 11.1% of visits record diabetes, "
     "making it second, whereas type 2 alone is 4.7%. Asthma (10.0%) and depression (9.5%) follow; "
     "obesity (3.6%) likely understates true prevalence because it depends on documentation. The share "
     "with a chronic condition rises from 13% of visits by children under 15 to 90% at 75 and older. "
     "These flags describe patients' conditions, not the reason for the visit, so they do not show whether "
     "visits were flare-ups or avoidable. They do show that ED clinicians routinely treat patients whose "
     "chronic illness can complicate care; whether better coordination would change ED use is a question "
     "these data cannot answer.")
para("**Injuries (Q7).** In all, 33.1% of visits (weighted) are coded as related to an injury (29.5%), "
     "an overdose or poisoning (1.4%) or an adverse effect of medical care (2.2%). Intent is recorded "
     "only for injuries and overdoses, and about a quarter of this combined group (25.7%, including every "
     "adverse-effect visit) has no intent recorded. Among visits with known intent, 7.3% are intentional "
     "and 92.7% unintentional. Overdoses and poisonings are more often recorded as intentional (13.2%), "
     "and among visits with known intent the intentional share is highest at ages 15-24 (12.1%).")
para("**Diagnostic services (Q8).** Imaging is used at 47.0% of all visits (weighted): X-ray 33.7%, CT "
     "16.5%, ultrasound 4.4% and MRI 0.7%. Use rises with age, from 29% of visits by children to 70% at 75 "
     "and older, and is higher for injury visits (54%) than other visits (44%). At least one blood test is "
     "ordered at 42.4% of visits. Among those visits, a complete blood count is by far the most common "
     "test (85.4%), followed by the comprehensive (55.3%) and basic (25.2%) metabolic panels, with 3.2 "
     "different tests per visit on average. Standalone glucose, BUN/creatinine and electrolyte orders "
     "are less common partly because the panels include those tests (codebook pp.3-4).")
para("**Medications (Q9).** The weighted mean is 2.49 medications per visit and the median is 2; 20.9% "
     "of visits involve none, 41.9% one or two and 37.2% three or more, matching NCHS (Table 1). On average "
     "1.62 are given in the ED and 1.08 prescribed at discharge (a drug can be both, so these overlap). The "
     "count rises with complexity, from 2.0 with no recorded chronic condition to 3.8 with three or more, "
     "and peaks at 3.1 for ages 45-64, so medication reconciliation and interaction checking are routine "
     "ED work.").paragraph_format.keep_together = True

# ---------------------------------------------------------------- 6. Q10
doc.add_heading("6. Correlations and outcomes (Q10)", level=1)
para("We built about 20 cleaned variables (missing codes removed, wait and length of visit winsorized) and "
     "computed weighted Pearson correlations, each using the visits where both variables are present. "
     "Triage level runs from 1 (immediate) to 5 (nonurgent), so a negative correlation means more of "
     "something among more urgent patients. The correlations describe associations, not causal effects.")
table([["Pair", "Weighted r", "Interpretation"],
       ["Age and Medicare as expected payer", "0.61", "Largely mechanical: Medicare eligibility at 65"],
       ["Age and number of chronic conditions", "0.55", "Chronic conditions accumulate with age"],
       ["Number of blood tests and admission", "0.37", "More extensive workups among admitted patients"],
       ["Triage level and number of blood tests", "-0.35", "More testing for more urgent patients"],
       ["Number of chronic conditions and admission", "0.33", "Patients with more conditions admitted more often"],
       ["Length of visit and number of blood tests", "0.31", "Longer visits involve more testing"],
       ["Wait and length of visit", "0.28", "Longer waits accompany longer visits"],
       ["Ambulance arrival and admission", "0.24", "Ambulance arrivals admitted more often"],
       ["Sex, ethnicity or pain score and any outcome", "about 0", "Little linear association"]],
      [2.85, 0.9, 3.15], "Table 3. Selected weighted correlations", left_cols=(0, 2))
figure([FIG + "q10_triage_admission_and_wait.png"], 5.3,
       "Figure 4. Weighted admission rate and median recorded wait by triage level (visits with triage "
       "levels 1-5; waits among visits with a recorded wait)")
para("Admission falls from 33% at triage level 1 to 2-3% at levels 4-5 (9.0% overall), consistent with "
     "triage identifying sicker patients. Recorded median waits are 14 minutes at level 1 but 18-20 "
     "minutes at every other level. That comparison is descriptive: wait time is blank for 15.2% of "
     "sampled visits and not applicable (not seen by a provider) for another 3.4%, so 18.6% have no "
     "usable wait; missingness differs by triage status and region; and case mix and hospital "
     "differences are not controlled. "
     "Whether emergent (level 2) patients wait longer than intended is a question for hospital-level "
     "timestamp data rather than something these medians establish. Deaths are too rare to analyze by "
     "subgroup: the sample contains 28 ED deaths and 7 patients dead on arrival, below the 30-record "
     "threshold NCHS uses for reliable estimates (codebook pp.8-9). We therefore report only the overall "
     "weighted rates (0.10% and 0.04%) and treat them as unreliable. Among admitted patients with a known "
     "discharge status, 40 of 1,752 died in hospital (2.6% weighted); with no standard errors, we treat "
     "this as descriptive only.")

# ---------------------------------------------------------------- 7. implications
doc.add_heading("7. Implications for hospital managers, and limitations", level=1)
bullet("Capacity: arrivals hold near peak from 10am to 8pm and run about 25% higher on Mondays than "
       "Sundays, and the median visit lasts 2.6 hours, so time in the department is as important a "
       "planning variable as arrival volume.")
bullet("Patient flow: recorded median waits for emergent (level 2) patients match those for nonurgent "
       "patients, a pattern worth checking against local timestamp data.")
bullet("Finance: public programs are the expected payer for about half of all visits; the effect on "
       "margins depends on reimbursement and collections not captured here.")
bullet("Limitations: no design-based standard errors; wait-time missingness varies by triage status "
       "and region; rare outcomes fall below NCHS's reliability threshold; condition flags depend on "
       "documentation; and 2015 form changes limit comparisons with earlier years.")
para("AI use. The course requires AI use. We used an AI coding assistant to write and "
     "debug the analysis code and to draft this report and the slides. We ran every step on the Yale "
     "HPC cluster, reviewed each output before continuing, and compared key estimates with figures "
     "published by NCHS, as described in Section 1.", size=9, color=GREY)

# ---------------------------------------------------------------- save (modern Word format)
compat = doc.settings.element.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat")
    doc.settings.element.append(compat)
cs = OxmlElement("w:compatSetting")
for k, v in [("name", "compatibilityMode"), ("uri", "http://schemas.microsoft.com/office/word"), ("val", "15")]:
    cs.set(qn("w:" + k), v)
compat.append(cs)
doc.save(OUT)
print("saved", OUT)
