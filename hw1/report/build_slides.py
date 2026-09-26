"""Build the HW1 slide deck (5 slides) from the verified results in hw1/results_log.md.

Charts are native PowerPoint charts built from logged values (no data re-run).
Run with:  uv run --with python-pptx python build_slides.py
"""
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HW1 = r"C:\Users\mnich\OneDrive\Documents\Second Year\EmpiricalStrategy\hw1"
OUT = sys.argv[1] if len(sys.argv) > 1 else HW1 + r"\report\MGT634_HW1_NHAMCS_slides.pptx"

NAVY = RGBColor(0x16, 0x32, 0x4F)     # dominant: title/closing backgrounds, headings
PANEL = RGBColor(0x22, 0x46, 0x6C)    # chart panels on dark slides
BLUE = RGBColor(0x2A, 0x78, 0xD6)     # data
CORAL = RGBColor(0xEB, 0x68, 0x34)    # accent: the one bar that carries the message
INK = RGBColor(0x1A, 0x1A, 0x1A)
SOFT = RGBColor(0x5A, 0x5F, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
GRIDLINE = RGBColor(0xE3, 0xE3, 0xE0)
PAYER_COLOURS = [BLUE, CORAL, RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0xED, 0xA1, 0x00),
                 RGBColor(0xE8, 0x7B, 0xA4)]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, colour):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = colour


def text(slide, x, y, w, h, content, size=16, colour=INK, bold=False, font="Calibri",
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=6):
    """Text box; `content` may be a list of (text, size, colour, bold) tuples, one per paragraph."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ["margin_left", "margin_right", "margin_top", "margin_bottom"]:
        setattr(tf, side, 0)
    items = content if isinstance(content, list) else [(content, size, colour, bold)]
    for i, (t, s, c, b) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = t
        r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(s), c, b, font
    return tb


def panel(slide, x, y, w, h, fill=PANEL):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False


def figure_stat(slide, x, y, w, number, label, num_colour=NAVY, label_colour=SOFT, size=38):
    """A large number with a short label underneath (no card behind it)."""
    text(slide, x, y, w, 0.75, number, size, num_colour, True, "Cambria")
    text(slide, x, y + 0.72, w, 0.9, label, 13, label_colour)


def title(slide, t, sub=None, colour=NAVY, sub_colour=SOFT):
    text(slide, 0.6, 0.35, 12.1, 1.0, t, 26, colour, True, "Cambria")
    if sub:
        text(slide, 0.6, 1.38, 12.1, 0.4, sub, 14, sub_colour)


def style_chart(chart, legend=False, font_size=11, colour=SOFT):
    chart.font.size, chart.font.name = Pt(font_size), "Calibri"
    chart.font.color.rgb = colour
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRIDLINE
    va.format.line.fill.background()
    chart.category_axis.format.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    chart.category_axis.has_major_gridlines = False


def label_font(dl, size, colour, bold=True):
    """Custom data-label text ignores dl.font - style the run itself."""
    r = dl.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
    r.font.color.rgb = colour


def value_labels(series, fmt, size=10, colour=SOFT, bold=False):
    dls = series.data_labels
    dls.show_value = True
    dls.number_format = fmt
    dls.number_format_is_linked = False
    dls.position = XL_LABEL_POSITION.OUTSIDE_END
    dls.font.size, dls.font.bold, dls.font.name = Pt(size), bold, "Calibri"
    dls.font.color.rgb = colour


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ============================================================ 1. scale + method
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.6, 0.55, 12, 0.4, "MGT 634 · HOMEWORK 1 · NHAMCS EMERGENCY DEPARTMENT DATA, 2015", 13, ICE, True)
text(s, 0.6, 0.95, 12.1, 1.4, "137 million ED visits: what the 2015 national sample shows hospital managers",
     34, WHITE, True, "Cambria")
text(s, 0.6, 2.2, 12, 0.4, "Michal Kozuchowski, Laila Lapins, Nick Giamalis, Raymond Chang, Sean Weller",
     14, ICE)
for i, (n, l) in enumerate([
        ("21,061", "sampled visits from 248 EDs; each row is one visit, not one patient"),
        ("136.9M", "U.S. ED visits once weighted by PATWT, matching the NCHS-published total"),
        ("-9 / -8 / -7", "codes NCHS uses for blank, unknown and not-applicable answers; recoded to missing")]):
    figure_stat(s, 0.6 + i * 4.1, 3.15, 3.7, n, l, num_colour=WHITE, label_colour=ICE, size=40)
text(s, 0.6, 5.25, 12.1, 1.6, [
    ("Method", 15, WHITE, True),
    ("Every estimate weighted by PATWT. Missing-value codes recoded before any statistic. Long waits and "
     "stays summarized with medians; means reported with values capped at the 99.5th percentile. Key "
     "estimates compared with figures NCHS publishes: all agree to rounding except wait-time "
     "missingness (15.2% vs 15.7%).", 14, ICE, False)])
notes(s, "Each row is one sampled ED visit (one Patient Record Form), not a patient. NCHS stores most "
         "missing answers as negative codes, so a standard NaN count understates missingness. Missing wait "
         "times are more common where triage was not recorded (38-45%) than at triage levels 1-5 (7-12%), "
         "and in the Northeast and West, so wait-time results describe visits with a recorded wait. All "
         "results are weighted point estimates without design-based standard errors.")

# ============================================================ 2. demand timing + throughput
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Arrivals stay near peak from 10am to 8pm, and the median visit lasts 2.6 hours",
      "Weighted share of daily arrivals by hour (visits with a recorded arrival time, 98.7% of records)")
cd = CategoryChartData()
cd.categories = [str(h) for h in range(24)]
hours = [2.20, 1.83, 1.62, 1.42, 1.36, 1.34, 1.76, 2.64, 3.59, 5.06, 6.05, 5.75, 6.08, 5.66,
         5.80, 5.61, 5.65, 5.77, 6.34, 5.90, 5.46, 5.00, 4.64, 3.46]
cd.add_series("% of daily arrivals", hours)
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.95), Inches(8.4),
                        Inches(5.1), cd).chart
style_chart(ch)
ch.plots[0].gap_width = 40
ser = ch.plots[0].series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = BLUE
pt = ser.points[18]                       # 6pm peak highlighted
pt.format.fill.solid()
pt.format.fill.fore_color.rgb = CORAL
pt.data_label.has_text_frame = True
pt.data_label.text_frame.text = "6pm: 6.3%"
label_font(pt.data_label, 12, CORAL)
pt.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.maximum_scale = 7.5
ch.value_axis.minimum_scale = 0
ch.value_axis.tick_labels.number_format = '0"%"'
ch.value_axis.tick_labels.number_format_is_linked = False
ch.category_axis.has_title = True
ch.category_axis.axis_title.text_frame.text = "Arrival hour (0 = midnight)"
label_font(ch.category_axis.axis_title, 11, SOFT, bold=False)
figure_stat(s, 9.4, 2.1, 3.4, "154 min", "median length of visit (mean 214 min)")
figure_stat(s, 9.4, 3.55, 3.4, "+25%", "visits per day on Mondays vs Sundays (432,000 vs 346,000)")
text(s, 9.4, 5.1, 3.4, 1.9, "Because patients stay about 2.6 hours, the number in the department "
     "likely peaks after arrivals do. Staffing plans should weigh time in the department, not arrivals "
     "alone.", 13, INK)
notes(s, "Weekdays are compared as average visits per day because 2015 had 53 Thursdays. Arrivals are "
         "lowest from 3am to 6am, climb from 7am, and stay near peak until about 8pm; 69% arrive between "
         "10am and 10pm, and the peak hour has 4.7 times the arrivals of the quietest. The Monday excess "
         "is consistent with demand deferred over the weekend, but the data cannot establish why. NHAMCS "
         "does not measure occupancy directly.")

# ============================================================ 3. payer mix
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Medicaid is the most common expected payer, and age largely determines who pays",
      "Primary expected payer by age group, weighted % of visits with a known expected payer (89.2% of visits)")
cd = CategoryChartData()
cd.categories = ["Under 15", "15-24", "25-44", "45-64", "65-74", "75+"]
mix = {"Medicaid/CHIP": [65.2, 40.8, 34.9, 26.0, 6.0, 3.3],
       "Private": [27.2, 37.7, 34.7, 41.4, 10.8, 7.9],
       "Medicare": [1.1, 2.5, 6.8, 18.0, 78.3, 86.9],
       "Self-pay": [4.7, 14.2, 17.5, 9.5, 2.3, 0.4],
       "Other": [1.9, 4.7, 6.1, 5.2, 2.6, 1.5]}
for k, v in mix.items():
    cd.add_series(k, v)
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, Inches(0.5), Inches(1.95), Inches(8.4),
                        Inches(5.2), cd).chart
style_chart(ch, legend=True)
ch.plots[0].gap_width = 45
ch.plots[0].overlap = 100
ch.category_axis.reverse_order = True     # youngest at the top
ch.value_axis.tick_labels.number_format = '0%'
ch.value_axis.tick_labels.number_format_is_linked = False
for ser, colour in zip(ch.plots[0].series, PAYER_COLOURS):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = colour
    ser.format.line.color.rgb = WHITE
for idx, name in [(0, "Medicaid/CHIP"), (2, "Medicare")]:   # label only the two public programs
    ser = ch.plots[0].series[idx]
    for j, v in enumerate(mix[name]):
        if v >= 5:                                              # skip thin slivers
            dl = ser.points[j].data_label
            dl.has_text_frame = True
            dl.text_frame.text = f"{v:.0f}%"
            label_font(dl, 11, WHITE)
            dl.position = XL_LABEL_POSITION.CENTER
figure_stat(s, 9.4, 2.1, 3.4, "48.9%", "of all visits list Medicaid or Medicare as primary expected payer "
            "(54.8% of visits with a known payer)")
figure_stat(s, 9.4, 3.8, 3.4, "11.0%", "self-pay or no charge (visits with a known payer)", num_colour=CORAL)
text(s, 9.4, 5.25, 3.4, 1.8, "Expected payer is not revenue. If public programs pay less per visit than "
     "private insurers, this mix limits per-visit revenue; the actual effect depends on contracts and "
     "collections.", 13, INK)
notes(s, "PAYTYPER is the primary expected source of payment, chosen by an NCHS hierarchy when several are "
         "listed; our all-visit shares match the NCHS table (Medicaid 31.15%). 10.8% of visits have a blank "
         "or unknown payer and are excluded from the chart. EDs must screen and stabilize patients "
         "regardless of ability to pay (EMTALA), so hospitals have little control over this mix at the door.")

# ============================================================ 4. clinical complexity
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Diabetes ranks second among chronic conditions once its three checkboxes are combined",
      "Weighted % of visits with the chronic-condition section completed (98.4% of records)")
cond = [("Hypertension", 23.97), ("Any diabetes (type 1, 2 or unspecified)", 11.06), ("Asthma", 9.98),
        ("Depression", 9.45), ("Hyperlipidemia", 8.17), ("Substance abuse", 6.65),
        ("Coronary artery disease", 6.08), ("COPD", 5.37)]
cd = CategoryChartData()
cd.categories = [c for c, _ in cond]
cd.add_series("% of visits", [v for _, v in cond])
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.95), Inches(8.4),
                        Inches(5.1), cd).chart
style_chart(ch)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 45
ser = ch.plots[0].series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = BLUE
value_labels(ser, '0.0"%"', size=11)
dpt = ser.points[1]                        # combined diabetes highlighted
dpt.format.fill.solid()
dpt.format.fill.fore_color.rgb = CORAL
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
figure_stat(s, 9.4, 2.1, 3.4, "4.7%", "type 2 diabetes alone; combining all three checkboxes gives 11.1%",
            num_colour=CORAL)
figure_stat(s, 9.4, 3.65, 3.4, "47.6%", "record at least one chronic condition (90% at ages 75+)")
figure_stat(s, 9.4, 5.2, 3.4, "37.2%", "of all visits involve three or more medications (median 2)")
notes(s, "The chronic-condition flags describe patients' conditions, not the reason for the visit, so they "
         "do not show whether visits were flare-ups or avoidable. Obesity (3.6%) likely understates "
         "prevalence because it depends on documentation. Medications rise from 2.0 per visit with no "
         "recorded chronic condition to 3.8 with three or more. Injury, overdose and adverse-effect visits "
         "(33.1% of visits) are covered in the write-up.")

# ============================================================ 5. triage
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
title(s, "Triage level tracks admission closely; recorded waits differ little beyond level 1",
      "Weighted, visits with triage levels 1-5 (1 = immediate, 5 = nonurgent)", colour=WHITE, sub_colour=ICE)
levels = ["1 Immediate", "2 Emergent", "3 Urgent", "4 Semi-urgent", "5 Nonurgent"]
for k, (vals, ttl, fmt, colour) in enumerate([
        ([33.34, 25.03, 11.40, 2.28, 2.53], "Admitted to hospital (%)", '0"%"', ICE),
        ([14, 18, 20, 19, 18], "Median recorded wait to see a provider (min)", "0", CORAL)]):
    x = 0.5 + k * 4.25
    panel(s, x, 1.95, 4.05, 3.5)
    text(s, x + 0.25, 2.08, 3.6, 0.4, ttl, 13, WHITE, True)
    cd = CategoryChartData()
    cd.categories = levels
    cd.add_series(ttl, vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x + 0.1), Inches(2.45),
                            Inches(3.85), Inches(2.9), cd).chart
    style_chart(ch, font_size=10, colour=ICE)
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.plots[0].gap_width = 50
    ser = ch.plots[0].series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = colour
    value_labels(ser, fmt, size=12, colour=WHITE, bold=True)
text(s, 9.15, 2.0, 3.65, 3.5, [
    ("Reading these results", 14, WHITE, True),
    ("Descriptive, not causal: case mix and hospital differences are not controlled.", 13, ICE, False),
    ("15% of visits lack a recorded wait, more often where triage was not recorded and in the Northeast "
     "and West.", 13, ICE, False),
    ("Weighted point estimates without survey-design standard errors.", 13, ICE, False)], after=9)
text(s, 0.6, 5.72, 12.1, 1.0, [
    ("A question for hospital managers", 15, WHITE, True),
    ("Emergent (level 2) and nonurgent (level 5) patients have the same recorded median wait, 18 minutes. "
     "Do local timestamp data show the same pattern, and if so, where does the delay occur?", 14, ICE, False)])
text(s, 0.6, 6.95, 12.1, 0.3, "AI use (course requirement): analysis code and slide drafts produced with an "
     "AI coding assistant; every result was run on the Yale cluster and reviewed by the group.", 10, ICE)
notes(s, "Admission falls from 33% at level 1 to 2-3% at levels 4-5 (9.0% overall), consistent with triage "
         "identifying sicker patients. The wait comparison is descriptive and uses only visits with a "
         "recorded wait. ED deaths (28 sampled records) and dead-on-arrival cases (7) are below the "
         "30-record threshold NCHS uses for reliable estimates, so they are not broken out here.")

prs.save(OUT)
print("saved", OUT)
